# -*- coding: utf-8 -*-
"""Built-in unit tests for Word Formatter Pro v2.7.4."""

from __future__ import annotations

import os
import tempfile
import unittest
import json
import base64
import tkinter as tk
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement

from wfp_config import DEFAULT_CONFIG
from wfp_core import (
    BLANK_LINE_MODE_DELETE_SINGLE,
    BLANK_LINE_MODE_KEEP_SINGLE,
    BLANK_LINE_MODE_PRESERVE,
    FormatReport,
    LegacyConversionUnavailable,
    WordProcessor,
    audit_caption_numbers,
    audit_cross_references,
    check_first_line_indent,
    check_caption_sequence,
    detect_inline_image_paragraph,
    detect_thesis_section,
    extract_figure_references_from_text,
    extract_table_references_from_text,
    find_nearby_figure_caption,
    find_nearby_table_caption,
    iter_block_items,
    is_body_paragraph,
    is_protected_thesis_paragraph,
    is_table_cell_body_paragraph,
    should_apply_first_line_indent,
    is_thesis_mode,
    parse_figure_caption_number,
    parse_table_caption_number,
    update_thesis_context,
)
from wfp_gui import (
    WordFormatterGUI,
    scale_geometry,
    scale_px,
    validate_ui_scale,
)


class TextNormalizationTests(unittest.TestCase):
    def test_symbol_normalization_keeps_decimal_numbers(self):
        self.assertEqual(
            WordProcessor._normalize_symbols_in_text("你好,世界."),
            "你好，世界。",
        )
        self.assertEqual(
            WordProcessor._normalize_symbols_in_text("3.14 是 pi"),
            "3.14 是 pi",
        )

    def test_ellipsis_and_quotes(self):
        self.assertEqual(
            WordProcessor._normalize_symbols_in_text('他说"你好"...'),
            "他说“你好”……",
        )
        self.assertEqual(
            WordProcessor._normalize_symbols_in_text("version 1.2"),
            "version 1.2",
        )

    def test_markdown_cleaning(self):
        raw = "# 标题\n**粗体** 和 [链接](https://example.com)\n![图片](a.png)\n> 引用"
        cleaned = WordProcessor._clean_markdown(raw)
        self.assertEqual(cleaned, "标题\n粗体 和 链接\n图片\n引用")

    def test_markdown_cleaning_preserves_source_numeric_numbering(self):
        raw = (
            "一、登录\n\n"
            "1. 打开软件\n"
            "2. 完成登录\n\n"
            "（一）首页搜索\n\n"
            "1. 输入关键词\n"
            "2. 点击院校\n"
            "1.2.3 小节号保持原样"
        )
        cleaned = WordProcessor._clean_markdown(raw)
        self.assertEqual(cleaned, raw)


class BlankLineTests(unittest.TestCase):
    def test_delete_single_blank_line_and_compress_multiple(self):
        text = "a\n\nb\n\n\nc"
        self.assertEqual(
            WordProcessor._remove_blank_lines_from_text(text),
            "a\nb\n\nc",
        )

    def test_keep_single_blank_line_and_compress_multiple(self):
        text = "a\n\nb\n\n\nc"
        self.assertEqual(
            WordProcessor._remove_blank_lines_from_text(text, keep_single_blank_lines=True),
            "a\n\nb\n\nc",
        )

    def test_blank_line_mode_aliases(self):
        self.assertEqual(
            WordProcessor._normalize_blank_line_mode("preserve"),
            BLANK_LINE_MODE_PRESERVE,
        )
        self.assertEqual(
            WordProcessor._normalize_blank_line_mode("delete_single"),
            BLANK_LINE_MODE_DELETE_SINGLE,
        )
        self.assertEqual(
            WordProcessor._normalize_blank_line_mode("keep_single"),
            BLANK_LINE_MODE_KEEP_SINGLE,
        )

    def test_processor_preserve_mode_leaves_text_unchanged(self):
        processor = WordProcessor({"blank_line_mode": BLANK_LINE_MODE_PRESERVE})
        text = "a\n\nb\n\n\nc"
        self.assertEqual(processor._normalize_text_blank_lines(text), text)


class TableHelperTests(unittest.TestCase):
    def test_numeric_table_text(self):
        self.assertTrue(WordProcessor._is_numeric_table_text("1,234.56"))
        self.assertTrue(WordProcessor._is_numeric_table_text("¥100元"))
        self.assertTrue(WordProcessor._is_numeric_table_text("12.5%"))
        self.assertFalse(WordProcessor._is_numeric_table_text("abc"))

    def test_short_table_text(self):
        self.assertTrue(WordProcessor._is_short_table_text("合计", max_len=4))
        self.assertFalse(WordProcessor._is_short_table_text("较长文本内容", max_len=4))

    def test_table_percentage_normalization(self):
        pcts = WordProcessor._normalize_table_pcts([1, 9], 20, 80)
        self.assertAlmostEqual(sum(pcts), 100.0)
        self.assertEqual(pcts, [20.0, 80.0])


class OoxmlProtectionTests(unittest.TestCase):
    def test_ooxml_element_detection(self):
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run()
        run._r.append(OxmlElement("w:drawing"))
        self.assertTrue(WordProcessor._has_drawing_or_pict(para))
        self.assertFalse(WordProcessor._has_field_codes(para))

        para_field = doc.add_paragraph()
        field_run = para_field.add_run()
        field_run._r.append(OxmlElement("w:fldChar"))
        self.assertTrue(WordProcessor._has_field_codes(para_field))

    def test_strip_leading_whitespace_removes_plain_blank_run(self):
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("   ")
        para.add_run("正文")
        WordProcessor({})._strip_leading_whitespace(para)
        self.assertEqual(para.text, "正文")

    def test_strip_leading_whitespace_preserves_special_run(self):
        doc = Document()
        para = doc.add_paragraph()
        special_run = para.add_run()
        special_run._r.append(OxmlElement("w:fldChar"))
        para.add_run("正文")

        WordProcessor({})._strip_leading_whitespace(para)

        self.assertTrue(WordProcessor._has_field_codes(para))
        self.assertEqual(len(para.runs), 2)
        self.assertEqual(para.text, "正文")


class FirstLineIndentTests(unittest.TestCase):
    def _format_doc(self, paragraphs, config=None, table_text=None):
        with tempfile.TemporaryDirectory(prefix="wfp_indent_test_") as tmpdir:
            source = Path(tmpdir) / "source.docx"
            output = Path(tmpdir) / "output.docx"
            doc = Document()
            for text in paragraphs:
                doc.add_paragraph(text)
            if table_text is not None:
                table = doc.add_table(rows=1, cols=1)
                table.cell(0, 0).text = table_text
            doc.save(source)

            merged_config = DEFAULT_CONFIG.copy()
            if config:
                merged_config.update(config)
            WordProcessor(merged_config).format_document(str(source), str(output))
            return Document(output)

    def test_body_paragraph_gets_first_line_indent(self):
        output = self._format_doc(["这是一个普通正文段落，用于验证首行缩进。"])

        self.assertTrue(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))

    def test_h1_paragraph_does_not_get_first_line_indent(self):
        output = self._format_doc(["一、绪论"])

        self.assertFalse(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))

    def test_h2_paragraph_does_not_get_first_line_indent(self):
        output = self._format_doc(["（一）研究背景"])

        self.assertFalse(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))

    def test_figure_caption_does_not_get_first_line_indent(self):
        output = self._format_doc(["图 1 系统结构图"])

        self.assertFalse(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))

    def test_table_caption_does_not_get_first_line_indent(self):
        output = self._format_doc(["表 1 实验结果"])

        self.assertFalse(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))

    def test_reference_entry_does_not_get_first_line_indent(self):
        output = self._format_doc(["[1] 张三. 文献标题. 期刊, 2024."])

        self.assertFalse(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))

    def test_table_cell_paragraph_is_not_body_paragraph(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        para = table.cell(0, 0).paragraphs[0]
        para.text = "表格内正文"

        self.assertFalse(is_body_paragraph(para))

    def test_old_config_without_first_line_indent_fields_is_compatible(self):
        old_config = DEFAULT_CONFIG.copy()
        for key in (
            "enable_first_line_indent",
            "first_line_indent_chars",
            "first_line_indent_tolerance_chars",
            "first_line_indent_scope",
        ):
            old_config.pop(key, None)

        output = self._format_doc(["旧配置正文段落。"], config=old_config)

        self.assertTrue(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))


class TableCellBodyIndentTests(unittest.TestCase):
    LONG_BODY = "本研究围绕开题报告表格模板中的正文段落展开，重点验证较长连续说明文字可以保留为正文并应用首行缩进。"

    def _format_table_doc(self, table_texts, config=None):
        with tempfile.TemporaryDirectory(prefix="wfp_table_body_indent_") as tmpdir:
            source = Path(tmpdir) / "source.docx"
            output = Path(tmpdir) / "output.docx"
            doc = Document()
            table = doc.add_table(rows=len(table_texts), cols=1)
            for row_idx, text in enumerate(table_texts):
                table.cell(row_idx, 0).text = text
            doc.save(source)

            merged_config = DEFAULT_CONFIG.copy()
            merged_config.update({
                "document_mode": "thesis",
                "thesis_mode_enabled": True,
                "enable_thesis_structure_detection": True,
                "protect_table_text": True,
            })
            if config:
                merged_config.update(config)
            report = WordProcessor(merged_config).format_document(str(source), str(output))
            return Document(output), report

    def _table_para(self, text):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        para = table.cell(0, 0).paragraphs[0]
        para.text = text
        return para

    def test_table_cell_short_labels_do_not_get_indent(self):
        output, report = self._format_table_doc(["题目", "学生姓名"])

        self.assertFalse(check_first_line_indent(output.tables[0].cell(0, 0).paragraphs[0], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.tables[0].cell(1, 0).paragraphs[0], 2.0, 0.2))
        self.assertGreaterEqual(report.table_cell_body_indent_skipped, 2)

    def test_table_cell_heading_does_not_get_indent(self):
        output, _ = self._format_table_doc(["一、选题的目的意义"])

        self.assertFalse(check_first_line_indent(output.tables[0].cell(0, 0).paragraphs[0], 2.0, 0.2))

    def test_table_cell_long_body_gets_indent_with_table_protection_enabled(self):
        output, report = self._format_table_doc([self.LONG_BODY])

        self.assertTrue(check_first_line_indent(output.tables[0].cell(0, 0).paragraphs[0], 2.0, 0.2))
        self.assertEqual(report.table_cell_body_paragraphs_detected, 1)
        self.assertEqual(report.table_cell_body_indent_fixed, 1)

    def test_table_cell_long_body_does_not_get_indent_when_disabled(self):
        output, report = self._format_table_doc(
            [self.LONG_BODY],
            config={"enable_table_cell_body_indent": False},
        )

        self.assertFalse(check_first_line_indent(output.tables[0].cell(0, 0).paragraphs[0], 2.0, 0.2))
        self.assertEqual(report.table_cell_body_paragraphs_detected, 0)

    def test_non_body_table_content_is_conservative(self):
        for text in (
            "图1 系统结构图",
            "表1 实验结果",
            "[1] 张三. 文献标题. 期刊, 2024.",
            "关键词：格式；论文；表格",
            "目录........................1",
        ):
            para = self._table_para(text)
            config = DEFAULT_CONFIG.copy()
            self.assertFalse(is_table_cell_body_paragraph(para, {"config": config}, config))
            self.assertFalse(should_apply_first_line_indent(para, {"config": config}, config))

    def test_old_config_without_table_cell_indent_fields_is_compatible(self):
        old_config = DEFAULT_CONFIG.copy()
        for key in (
            "enable_table_cell_body_indent",
            "table_cell_body_indent_min_chars",
            "table_cell_body_indent_scope",
        ):
            old_config.pop(key, None)

        output, _ = self._format_table_doc([self.LONG_BODY], config=old_config)

        self.assertTrue(check_first_line_indent(output.tables[0].cell(0, 0).paragraphs[0], 2.0, 0.2))

    def test_format_report_counts_table_cell_body_indent(self):
        _, report = self._format_table_doc(["题目", "一、选题的目的意义", self.LONG_BODY])

        self.assertEqual(report.table_cell_body_paragraphs_detected, 1)
        self.assertEqual(report.table_cell_body_indent_fixed, 1)
        self.assertGreaterEqual(report.table_cell_body_indent_skipped, 2)
        self.assertIn("检测到表格内长正文段落数量：1", report.to_text())


class ThesisModeDetectionTests(unittest.TestCase):
    def _paragraph(self, text):
        doc = Document()
        return doc.add_paragraph(text)

    def test_thesis_config_defaults_and_old_config_compatibility(self):
        self.assertEqual(DEFAULT_CONFIG["document_mode"], "general")
        self.assertFalse(DEFAULT_CONFIG["thesis_mode_enabled"])
        self.assertTrue(DEFAULT_CONFIG["enable_thesis_structure_detection"])
        self.assertFalse(is_thesis_mode(DEFAULT_CONFIG.copy()))
        self.assertTrue(is_thesis_mode({**DEFAULT_CONFIG, "document_mode": "thesis"}))
        self.assertTrue(is_thesis_mode({**DEFAULT_CONFIG, "thesis_mode_enabled": True}))

        old_config = DEFAULT_CONFIG.copy()
        for key in (
            "document_mode",
            "thesis_mode_enabled",
            "protect_toc",
            "protect_references",
            "protect_captions",
            "protect_equations",
            "protect_table_text",
            "enable_thesis_structure_detection",
        ):
            old_config.pop(key, None)
        self.assertFalse(is_thesis_mode(old_config))

    def test_detects_core_thesis_sections(self):
        cases = [
            ("摘要", "chinese_abstract"),
            ("摘 要", "chinese_abstract"),
            ("摘要：本文研究投饵控制系统。", "chinese_abstract"),
            ("ABSTRACT", "english_abstract"),
            ("Abstract", "english_abstract"),
            ("英文摘要", "english_abstract"),
            ("关键词：PLC；投饵；排污", "keywords"),
            ("Keywords: PLC; feeding; drainage", "keywords"),
            ("目录", "toc"),
            ("绪论................1", "toc"),
            ("第1章 绪论", "heading"),
            ("第一章 绪论", "heading"),
            ("1 绪论", "heading"),
            ("1.1 研究背景", "heading"),
            ("图1-1 系统结构图", "figure_caption"),
            ("图 1.1 系统结构图", "figure_caption"),
            ("Figure 1 System architecture", "figure_caption"),
            ("表1-1 I/O分配表", "table_caption"),
            ("表 1.1 I/O分配表", "table_caption"),
            ("Table 1 I/O allocation", "table_caption"),
            ("（1-1）", "equation"),
            ("参考文献", "references_heading"),
            ("[1] 张三. 标题...", "reference_item"),
            ("1. 张三. 标题...", "reference_item"),
            ("致谢", "acknowledgement"),
            ("谢辞", "acknowledgement"),
            ("附录A", "appendix"),
            ("Appendix A", "appendix"),
        ]
        for text, expected in cases:
            with self.subTest(text=text):
                self.assertEqual(detect_thesis_section(self._paragraph(text)), expected)

    def test_context_state_machine_counts_and_protects(self):
        config = {**DEFAULT_CONFIG, "document_mode": "thesis"}
        context = {}
        for text in ["目录", "绪论................1", "第1章 绪论", "正文段落", "参考文献", "[1] 张三. 标题...", "附录A"]:
            para = self._paragraph(text)
            update_thesis_context(para, context)

        self.assertFalse(context["in_toc"])
        self.assertTrue(context["in_references"])
        self.assertTrue(context["in_appendix"])
        self.assertEqual(context["heading_count"], 1)
        self.assertEqual(context["toc_paragraph_count"], 2)
        self.assertEqual(context["reference_item_count"], 1)

        toc_context = {}
        toc_para = self._paragraph("目录")
        update_thesis_context(toc_para, toc_context)
        self.assertTrue(is_protected_thesis_paragraph(toc_para, toc_context, config))

        ref_context = {"in_references": True}
        self.assertTrue(is_protected_thesis_paragraph(self._paragraph("[1] 张三. 标题..."), ref_context, config))
        self.assertTrue(is_protected_thesis_paragraph(self._paragraph("图1-1 系统结构图"), {}, config))
        self.assertTrue(is_protected_thesis_paragraph(self._paragraph("表1-1 I/O分配表"), {}, config))

        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table_para = table.cell(0, 0).paragraphs[0]
        table_para.text = "表格内文字"
        self.assertTrue(is_protected_thesis_paragraph(table_para, {}, config))


class ThesisModeFormattingTests(unittest.TestCase):
    def _format_doc(self, paragraphs, config=None, table_text=None):
        with tempfile.TemporaryDirectory(prefix="wfp_thesis_test_") as tmpdir:
            source = Path(tmpdir) / "source.docx"
            output = Path(tmpdir) / "output.docx"
            doc = Document()
            for text in paragraphs:
                doc.add_paragraph(text)
            if table_text is not None:
                table = doc.add_table(rows=1, cols=1)
                table.cell(0, 0).text = table_text
            doc.save(source)

            merged_config = DEFAULT_CONFIG.copy()
            merged_config.update({
                "document_mode": "thesis",
                "thesis_mode_enabled": True,
                "enable_thesis_structure_detection": True,
            })
            if config:
                merged_config.update(config)
            report = WordProcessor(merged_config).format_document(str(source), str(output))
            return Document(output), report

    def test_thesis_mode_protects_non_body_and_indents_body(self):
        paragraphs = [
            "目录",
            "第1章 绪论",
            "这是一个普通正文段落，用于验证论文模式仍然应用正文首行缩进。",
            "图1-1 系统结构图",
            "表1-1 I/O分配表",
            "参考文献",
            "[1] 张三. 标题...",
            "（1-1）",
        ]

        output, report = self._format_doc(paragraphs, table_text="表格内文字")

        self.assertFalse(check_first_line_indent(output.paragraphs[0], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.paragraphs[1], 2.0, 0.2))
        self.assertTrue(check_first_line_indent(output.paragraphs[2], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.paragraphs[3], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.paragraphs[4], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.paragraphs[6], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.paragraphs[7], 2.0, 0.2))
        self.assertFalse(check_first_line_indent(output.tables[0].cell(0, 0).paragraphs[0], 2.0, 0.2))

        self.assertTrue(report.thesis_mode_enabled)
        self.assertGreaterEqual(report.toc_paragraphs_detected, 1)
        self.assertGreaterEqual(report.reference_items_detected, 1)
        self.assertGreaterEqual(report.figure_captions_detected, 1)
        self.assertGreaterEqual(report.table_captions_detected, 1)
        self.assertGreaterEqual(report.equations_detected, 1)
        self.assertGreaterEqual(report.protected_paragraphs, 6)
        self.assertIn("论文结构识别摘要", report.to_text())


class CaptionAuditTests(unittest.TestCase):
    TINY_PNG = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
    )

    def _paragraph(self, text):
        doc = Document()
        return doc.add_paragraph(text)

    def _image_doc(self, trailing_paragraphs):
        doc = Document()
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as image_file:
            image_file.write(self.TINY_PNG)
            image_path = image_file.name
        try:
            para = doc.add_paragraph()
            para.add_run().add_picture(image_path)
        finally:
            os.remove(image_path)
        for text in trailing_paragraphs:
            doc.add_paragraph(text)
        return doc

    def test_figure_caption_patterns_and_body_references_are_distinct(self):
        valid = [
            "图1 系统结构图",
            "图1-1 系统结构图",
            "图 1-1 系统结构图",
            "图1.1 系统结构图",
            "图1-1(a) 系统结构图",
            "图1-1（a）系统结构图",
            "Figure 1 System Architecture",
            "Fig. 1 System Architecture",
        ]
        for text in valid:
            with self.subTest(text=text):
                self.assertIsNotNone(parse_figure_caption_number(text))
        self.assertIsNone(parse_figure_caption_number("如图1-1所示"))

    def test_table_caption_patterns_and_body_references_are_distinct(self):
        valid = [
            "表1 I/O分配表",
            "表1-1 I/O分配表",
            "表 1-1 I/O分配表",
            "表1.1 I/O分配表",
            "Table 1 I/O Allocation",
        ]
        for text in valid:
            with self.subTest(text=text):
                self.assertIsNotNone(parse_table_caption_number(text))
        self.assertIsNone(parse_table_caption_number("见表1-1"))

    def test_caption_sequence_detects_duplicates_skips_and_subfigure_rules(self):
        figure_numbers = [
            parse_figure_caption_number("图1-1 系统结构图"),
            parse_figure_caption_number("图1-1 另一标题"),
            parse_figure_caption_number("图1-3 控制流程图"),
            parse_figure_caption_number("图1-1(a) 子图A"),
            parse_figure_caption_number("图1-1(b) 子图B"),
            parse_figure_caption_number("图1-1(a) 子图A重复"),
        ]
        result = check_caption_sequence(figure_numbers, "图")
        self.assertIn("1-1", result["duplicate_numbers"])
        self.assertIn("1-2", result["skipped_numbers"])
        self.assertNotIn("1-1(a)", result["duplicate_numbers"])
        self.assertIn("1-1(a)", result["duplicate_subfigure_numbers"])

        table_numbers = [
            parse_table_caption_number("表1-1 I/O分配表"),
            parse_table_caption_number("表1-1 另一标题"),
            parse_table_caption_number("表1-3 控制点表"),
        ]
        table_result = check_caption_sequence(table_numbers, "表")
        self.assertIn("1-1", table_result["duplicate_numbers"])
        self.assertIn("1-2", table_result["skipped_numbers"])

    def test_position_helpers_find_nearby_captions_and_ignore_missing_ones(self):
        doc = self._image_doc(["", "图1-1 系统结构图"])
        blocks = list(iter_block_items(doc))
        self.assertTrue(detect_inline_image_paragraph(blocks[0]))
        self.assertIsNotNone(find_nearby_figure_caption(blocks, 0, window=3))

        doc_without_caption = self._image_doc(["正文说明", "继续说明", "仍然没有题注"])
        blocks_without_caption = list(iter_block_items(doc_without_caption))
        self.assertIsNone(find_nearby_figure_caption(blocks_without_caption, 0, window=3))

        table_doc = Document()
        table_doc.add_paragraph("表1-1 I/O分配表")
        table_doc.add_table(rows=1, cols=1)
        table_blocks = list(iter_block_items(table_doc))
        self.assertIsNotNone(find_nearby_table_caption(table_blocks, 1, window=3))

        table_doc_without_caption = Document()
        table_doc_without_caption.add_table(rows=1, cols=1)
        table_doc_without_caption.add_paragraph("正文说明")
        missing_blocks = list(iter_block_items(table_doc_without_caption))
        self.assertIsNone(find_nearby_table_caption(missing_blocks, 0, window=3))

    def test_cross_reference_audit_respects_existing_and_missing_numbers(self):
        doc = Document()
        doc.add_paragraph("图1-1 系统结构图")
        doc.add_paragraph("表1-1 I/O分配表")
        doc.add_paragraph("如图1-1所示，系统结构清晰。")
        doc.add_paragraph("如图1-9所示，备用结构如下。")
        doc.add_paragraph("见表1-1。")
        doc.add_paragraph("见表1-9。")
        report = FormatReport()
        audit_caption_numbers(doc, report, DEFAULT_CONFIG.copy())
        audit_cross_references(doc, report, DEFAULT_CONFIG.copy())

        self.assertGreaterEqual(report.figure_references_detected, 2)
        self.assertGreaterEqual(report.table_references_detected, 2)
        self.assertIn("1-9", report.missing_figure_references)
        self.assertIn("1-9", report.missing_table_references)

    def test_unreferenced_caption_warning_is_config_controlled(self):
        doc = Document()
        doc.add_paragraph("图1-1 系统结构图")
        report = FormatReport()
        audit_caption_numbers(doc, report, DEFAULT_CONFIG.copy())
        audit_cross_references(doc, report, {**DEFAULT_CONFIG, "warn_unreferenced_captions": False})
        self.assertFalse(any("未被正文引用" in warning for warning in report.warnings))

        report_enabled = FormatReport()
        audit_caption_numbers(doc, report_enabled, DEFAULT_CONFIG.copy())
        audit_cross_references(
            doc,
            report_enabled,
            {**DEFAULT_CONFIG, "warn_unreferenced_captions": True},
        )
        self.assertTrue(any("未被正文引用" in warning for warning in report_enabled.warnings))

    def test_config_toggles_window_fallback_and_report_text(self):
        old_config = DEFAULT_CONFIG.copy()
        for key in (
            "enable_caption_number_audit",
            "enable_cross_reference_audit",
            "caption_search_window",
            "caption_numbering_mode",
            "allow_subfigure_suffix",
            "warn_unreferenced_captions",
        ):
            old_config.pop(key, None)

        doc = Document()
        doc.add_paragraph("图1-1 系统结构图")
        doc.add_paragraph("表1-1 I/O分配表")

        disabled_report = FormatReport()
        audit_caption_numbers(doc, disabled_report, {**DEFAULT_CONFIG, "enable_caption_number_audit": False})
        self.assertEqual(disabled_report.figure_captions_detected, 0)

        report = FormatReport()
        audit_caption_numbers(doc, report, {**old_config, "caption_search_window": 99})
        audit_cross_references(doc, report, {**old_config, "enable_cross_reference_audit": False})
        text = report.to_text()
        self.assertEqual(report.caption_search_window_used, 3)
        self.assertIn("图题检查摘要", text)
        self.assertIn("表题检查摘要", text)
        self.assertIn("未发现明显图表题注编号问题", text)

    def test_reference_extractors(self):
        self.assertIn("1-1", extract_figure_references_from_text("如图1-1所示"))
        self.assertIn("1-1", extract_figure_references_from_text("见图 1-1"))
        self.assertIn("1.1", extract_figure_references_from_text("由图1.1可知"))
        self.assertIn("1", extract_figure_references_from_text("Figure 1 shows the system"))
        self.assertIn("1-1", extract_table_references_from_text("如表1-1所示"))
        self.assertIn("1-1", extract_table_references_from_text("见表 1-1"))
        self.assertIn("1", extract_table_references_from_text("Table 1 lists the signals"))


class FormatReportTests(unittest.TestCase):
    def test_to_text_outputs_key_fields(self):
        report = FormatReport(
            input_file="input.docx",
            output_file="output_formatted.docx",
            total_paragraphs=3,
            total_tables=1,
            body_paragraphs=2,
            headings_detected=1,
            captions_detected=1,
            first_line_indent_fixed=2,
            headings_fixed=1,
            tables_fixed=1,
            captions_fixed=1,
            skipped_toc_paragraphs=1,
            skipped_reference_paragraphs=1,
            skipped_table_paragraphs=2,
            report_file="output_format_report.txt",
        )

        text = report.to_text()

        self.assertIn("输入文件：input.docx", text)
        self.assertIn("输出文件：output_formatted.docx", text)
        self.assertIn("处理时间：", text)
        self.assertIn("总段落数：3", text)
        self.assertIn("修复首行缩进：2", text)
        self.assertIn("报告文件：output_format_report.txt", text)
        self.assertIn("Warnings：无", text)
        self.assertIn("Errors：无", text)

    def test_add_warning_and_add_error(self):
        report = FormatReport()

        report.add_warning("无法准确统计某字段")
        report.add_error("报告写入失败")

        self.assertEqual(report.warnings, ["无法准确统计某字段"])
        self.assertEqual(report.errors, ["报告写入失败"])
        text = report.to_text()
        self.assertIn("- 无法准确统计某字段", text)
        self.assertIn("- 报告写入失败", text)


class FormatReportFileTests(unittest.TestCase):
    def _write_source_doc(self, source):
        doc = Document()
        doc.add_paragraph("一、绪论")
        doc.add_paragraph("这是一个普通正文段落，用于验证报告生成。")
        doc.add_paragraph("目录")
        doc.add_paragraph("第一章 绪论........1")
        doc.add_paragraph("[1] 张三. 文献标题. 期刊, 2024.")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格内正文"
        doc.save(source)

    def test_report_file_is_generated_with_utf8(self):
        with tempfile.TemporaryDirectory(prefix="wfp_report_test_") as tmpdir:
            source = Path(tmpdir) / "source.docx"
            output = Path(tmpdir) / "source_formatted.docx"
            self._write_source_doc(source)

            report = WordProcessor(DEFAULT_CONFIG.copy()).format_document(str(source), str(output))
            report_path = Path(report.report_file)

            self.assertTrue(output.exists())
            self.assertTrue(report_path.exists())
            self.assertEqual(report_path.name, "source_format_report.txt")
            text = report_path.read_text(encoding="utf-8")
            self.assertIn("输入文件：", text)
            self.assertIn("统计摘要", text)
            self.assertIn("总段落数：", text)

    def test_report_file_is_not_generated_when_disabled(self):
        with tempfile.TemporaryDirectory(prefix="wfp_report_disabled_") as tmpdir:
            source = Path(tmpdir) / "source.docx"
            output = Path(tmpdir) / "source_formatted.docx"
            self._write_source_doc(source)

            config = DEFAULT_CONFIG.copy()
            config["enable_format_report"] = False
            report = WordProcessor(config).format_document(str(source), str(output))

            self.assertTrue(output.exists())
            self.assertIsNone(report.report_file)
            self.assertFalse((Path(tmpdir) / "source_format_report.txt").exists())

    def test_old_config_without_report_fields_is_compatible(self):
        with tempfile.TemporaryDirectory(prefix="wfp_old_report_config_") as tmpdir:
            source = Path(tmpdir) / "source.docx"
            output = Path(tmpdir) / "source_formatted.docx"
            self._write_source_doc(source)

            old_config = DEFAULT_CONFIG.copy()
            old_config.pop("enable_format_report", None)
            old_config.pop("report_level", None)

            report = WordProcessor(old_config).format_document(str(source), str(output))

            self.assertTrue(output.exists())
            self.assertTrue(Path(report.report_file).exists())


class GUIFirstLineIndentConfigTests(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.message_patches = [
            mock.patch("wfp_gui.messagebox.showinfo"),
            mock.patch("wfp_gui.messagebox.showerror"),
            mock.patch("wfp_gui.messagebox.showwarning"),
        ]
        for patcher in self.message_patches:
            patcher.start()
        self.app = WordFormatterGUI(self.root)

    def tearDown(self):
        for patcher in reversed(self.message_patches):
            patcher.stop()
        self.root.destroy()

    def _entry_value(self, key):
        return self.app.entries[key].get().strip()

    def test_gui_first_line_indent_defaults_initialize(self):
        self.assertTrue(self.app.enable_first_line_indent_var.get())
        self.assertEqual(self._entry_value("first_line_indent_chars"), "2.0")
        self.assertEqual(self._entry_value("first_line_indent_tolerance_chars"), "0.2")
        self.assertEqual(self._entry_value("first_line_indent_scope"), "body_only")

    def test_save_config_includes_first_line_indent_fields(self):
        with tempfile.TemporaryDirectory(prefix="wfp_gui_config_") as tmpdir:
            output = Path(tmpdir) / "config.json"
            with mock.patch("wfp_gui.filedialog.asksaveasfilename", return_value=str(output)):
                self.app.save_config()

            saved = json.loads(output.read_text(encoding="utf-8"))

        self.assertTrue(saved["enable_first_line_indent"])
        self.assertEqual(saved["first_line_indent_chars"], 2.0)
        self.assertEqual(saved["first_line_indent_tolerance_chars"], 0.2)
        self.assertEqual(saved["first_line_indent_scope"], "body_only")
        self.assertTrue(saved["enable_table_cell_body_indent"])
        self.assertEqual(saved["table_cell_body_indent_min_chars"], 25)
        self.assertTrue(saved["enable_format_report"])
        self.assertEqual(saved["report_level"], "normal")

    def test_gui_thesis_mode_defaults_and_save_config(self):
        self.assertEqual(self.app.document_mode_var.get(), "普通文档")
        self.assertTrue(self.app.enable_thesis_structure_detection_var.get())
        self.assertTrue(self.app.protect_toc_var.get())
        self.assertTrue(self.app.protect_references_var.get())
        self.assertTrue(self.app.protect_captions_var.get())
        self.assertTrue(self.app.protect_equations_var.get())
        self.assertTrue(self.app.protect_table_text_var.get())

        self.app.document_mode_var.set("论文模式")
        self.app.protect_equations_var.set(False)
        config = self.app.collect_config()

        self.assertEqual(config["document_mode"], "thesis")
        self.assertTrue(config["thesis_mode_enabled"])
        self.assertTrue(config["enable_thesis_structure_detection"])
        self.assertTrue(config["protect_toc"])
        self.assertTrue(config["protect_references"])
        self.assertTrue(config["protect_captions"])
        self.assertFalse(config["protect_equations"])
        self.assertTrue(config["protect_table_text"])

    def test_old_config_load_uses_first_line_indent_defaults(self):
        old_config = DEFAULT_CONFIG.copy()
        for key in (
            "enable_first_line_indent",
            "first_line_indent_chars",
            "first_line_indent_tolerance_chars",
            "first_line_indent_scope",
            "enable_table_cell_body_indent",
            "table_cell_body_indent_min_chars",
            "table_cell_body_indent_scope",
        ):
            old_config.pop(key, None)

        self.app._apply_config(old_config)
        collected = self.app.collect_config()

        self.assertTrue(collected["enable_first_line_indent"])
        self.assertEqual(collected["first_line_indent_chars"], 2.0)
        self.assertEqual(collected["first_line_indent_tolerance_chars"], 0.2)
        self.assertEqual(collected["first_line_indent_scope"], "body_only")
        self.assertTrue(collected["enable_table_cell_body_indent"])
        self.assertEqual(collected["table_cell_body_indent_min_chars"], 25)
        self.assertEqual(collected["table_cell_body_indent_scope"], "long_text_only")
        self.assertTrue(collected["enable_format_report"])
        self.assertEqual(collected["report_level"], "normal")

    def test_old_config_load_uses_thesis_mode_defaults(self):
        old_config = DEFAULT_CONFIG.copy()
        for key in (
            "document_mode",
            "thesis_mode_enabled",
            "protect_toc",
            "protect_references",
            "protect_captions",
            "protect_equations",
            "protect_table_text",
            "enable_thesis_structure_detection",
        ):
            old_config.pop(key, None)

        self.app._apply_config(old_config)
        collected = self.app.collect_config()

        self.assertEqual(collected["document_mode"], "general")
        self.assertFalse(collected["thesis_mode_enabled"])
        self.assertTrue(collected["enable_thesis_structure_detection"])
        self.assertTrue(collected["protect_toc"])
        self.assertTrue(collected["protect_references"])
        self.assertTrue(collected["protect_captions"])
        self.assertTrue(collected["protect_equations"])
        self.assertTrue(collected["protect_table_text"])

    def test_invalid_first_line_indent_chars_is_blocked_on_save(self):
        self.app.entries["first_line_indent_chars"].delete(0, tk.END)
        self.app.entries["first_line_indent_chars"].insert(0, "5.1")

        with tempfile.TemporaryDirectory(prefix="wfp_gui_config_") as tmpdir:
            output = Path(tmpdir) / "config.json"
            with mock.patch("wfp_gui.filedialog.asksaveasfilename", return_value=str(output)):
                self.app.save_config()

            self.assertFalse(output.exists())

    def test_invalid_first_line_indent_tolerance_is_blocked_on_save(self):
        self.app.entries["first_line_indent_tolerance_chars"].delete(0, tk.END)
        self.app.entries["first_line_indent_tolerance_chars"].insert(0, "1.1")

        with tempfile.TemporaryDirectory(prefix="wfp_gui_config_") as tmpdir:
            output = Path(tmpdir) / "config.json"
            with mock.patch("wfp_gui.filedialog.asksaveasfilename", return_value=str(output)):
                self.app.save_config()

        self.assertFalse(output.exists())

    def test_invalid_table_cell_body_min_chars_falls_back_on_save(self):
        self.app.entries["table_cell_body_indent_min_chars"].delete(0, tk.END)
        self.app.entries["table_cell_body_indent_min_chars"].insert(0, "bad")

        config = self.app.collect_config()

        self.assertTrue(self.app.validate_config(config))
        self.assertEqual(config["table_cell_body_indent_min_chars"], 25)
        self.assertEqual(config["table_cell_body_indent_scope"], "long_text_only")


class GUIUiScaleConfigTests(unittest.TestCase):
    def setUp(self):
        self.root = tk.Tk()
        self.root.withdraw()
        self.message_patches = [
            mock.patch("wfp_gui.messagebox.showinfo"),
            mock.patch("wfp_gui.messagebox.showerror"),
            mock.patch("wfp_gui.messagebox.showwarning"),
        ]
        for patcher in self.message_patches:
            patcher.start()
        self.app = WordFormatterGUI(self.root)

    def tearDown(self):
        for patcher in reversed(self.message_patches):
            patcher.stop()
        self.root.destroy()

    def test_ui_scale_defaults_initialize(self):
        self.assertEqual(DEFAULT_CONFIG["ui_scale"], 1.0)
        self.assertTrue(DEFAULT_CONFIG["remember_window_geometry"])
        self.assertEqual(self.app.ui_scale_var.get(), "100%")
        self.assertTrue(self.app.remember_window_geometry_var.get())

    def test_old_config_load_uses_ui_scale_defaults(self):
        old_config = DEFAULT_CONFIG.copy()
        old_config.pop("ui_scale", None)
        old_config.pop("remember_window_geometry", None)

        self.app._apply_config(old_config)
        collected = self.app.collect_config()

        self.assertEqual(collected["ui_scale"], 1.0)
        self.assertTrue(collected["remember_window_geometry"])

    def test_invalid_ui_scale_falls_back_to_default(self):
        self.assertEqual(validate_ui_scale("2.0"), 1.0)
        self.app._apply_config({**DEFAULT_CONFIG, "ui_scale": 2.0})

        self.assertEqual(self.app.collect_config()["ui_scale"], 1.0)
        self.assertEqual(self.app.ui_scale_var.get(), "100%")

    def test_scale_helpers_return_reasonable_sizes(self):
        self.assertEqual(scale_px(100, 1.25), 125)
        self.assertEqual(scale_geometry(1200, 860, 1.25), (1500, 1075))

    def test_save_config_includes_ui_scale(self):
        with tempfile.TemporaryDirectory(prefix="wfp_gui_config_") as tmpdir:
            output = Path(tmpdir) / "config.json"
            self.app.ui_scale_var.set("125%")
            self.app.remember_window_geometry_var.set(False)
            with mock.patch("wfp_gui.filedialog.asksaveasfilename", return_value=str(output)):
                self.app.save_config()

            saved = json.loads(output.read_text(encoding="utf-8"))

        self.assertEqual(saved["ui_scale"], 1.25)
        self.assertFalse(saved["remember_window_geometry"])

    def test_load_config_restores_ui_scale(self):
        self.app._apply_config({**DEFAULT_CONFIG, "ui_scale": 1.5, "remember_window_geometry": False})

        self.assertEqual(self.app.collect_config()["ui_scale"], 1.5)
        self.assertEqual(self.app.ui_scale_var.get(), "150%")
        self.assertFalse(self.app.remember_window_geometry_var.get())


class UiScaleFormattingIsolationTests(unittest.TestCase):
    def test_ui_scale_does_not_change_word_formatting_defaults(self):
        baseline = DEFAULT_CONFIG.copy()
        scaled = DEFAULT_CONFIG.copy()
        scaled["ui_scale"] = 1.5

        word_formatting_keys = [
            "title_size",
            "h1_size",
            "h2_size",
            "body_size",
            "page_number_size",
            "table_caption_size",
            "figure_caption_size",
            "attachment_size",
            "subtitle_size",
            "margin_top",
            "margin_bottom",
            "margin_left",
            "margin_right",
            "line_spacing",
            "first_line_indent_chars",
            "table_size",
            "table_line_spacing",
        ]

        for key in word_formatting_keys:
            self.assertEqual(scaled[key], baseline[key])


class TempAndConversionTests(unittest.TestCase):
    def test_temp_docx_path_uses_system_temp_and_safe_name(self):
        processor = WordProcessor(DEFAULT_CONFIG.copy())
        temp_path = processor._make_temp_docx_path("copy", 'bad<name>:"?')
        try:
            self.assertEqual(
                os.path.normcase(os.path.abspath(os.path.dirname(temp_path))),
                os.path.normcase(os.path.abspath(tempfile.gettempdir())),
            )
            self.assertRegex(os.path.basename(temp_path), r"^~temp_copy_bad_name_[0-9]+_[0-9a-f]{8}\.docx$")
            self.assertIn(temp_path, processor.temp_files)
        finally:
            processor._cleanup_temp_files()

    def test_txt_conversion_uses_blank_line_mode_and_cleans_temp_file(self):
        with tempfile.TemporaryDirectory(prefix="wfp273_test_") as tmpdir:
            source = Path(tmpdir) / "sample.txt"
            source.write_text("标题\n\n正文一\n\n\n正文二", encoding="utf-8")

            processor = WordProcessor(DEFAULT_CONFIG.copy())
            temp_docx, is_from_txt = processor.convert_to_docx(str(source))
            try:
                self.assertTrue(is_from_txt)
                self.assertTrue(os.path.exists(temp_docx))
                self.assertEqual(
                    os.path.normcase(os.path.abspath(os.path.dirname(temp_docx))),
                    os.path.normcase(os.path.abspath(tempfile.gettempdir())),
                )
                self.assertFalse(list(Path(tmpdir).glob("~temp_*.docx")))

                converted = Document(temp_docx)
                self.assertEqual(
                    [p.text for p in converted.paragraphs],
                    ["标题", "正文一", "", "正文二"],
                )
            finally:
                processor._cleanup_temp_files()
            self.assertFalse(os.path.exists(temp_docx))

    def test_missing_soffice_marks_legacy_conversion_skipped(self):
        class MissingConverter:
            available = False

        processor = WordProcessor(DEFAULT_CONFIG.copy())
        processor.soffice_converter = MissingConverter()
        with self.assertRaises(LegacyConversionUnavailable):
            processor._convert_legacy_with_soffice("legacy.doc", "unused.docx")


def main(argv=None):
    suite = unittest.defaultTestLoader.loadTestsFromModule(__import__(__name__))
    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)
    if result.wasSuccessful():
        print("所有单元测试通过")
        return 0
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
